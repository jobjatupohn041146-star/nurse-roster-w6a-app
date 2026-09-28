"""Shared helpers: read the roster workbook, normalise codes, validate rules."""
import re
import openpyxl
from openpyxl.utils import get_column_letter

DAY_COL0 = 6          # column F = day 1 (template layout)
HEADER_ROW = 2
FIRST_NURSE_ROW = 4

WORK = {"C4", "P4", "CK", "TRAIN"}
LEAVE = {"V", "HBD", "LK", "LP", "H"}        # justified leave codes
PLAIN_OFF = {"X", "-"}                       # weekly/public holiday, or plain off
OFF = LEAVE | PLAIN_OFF
DAY_SHIFTS = {"C4", "CK"}                    # start 07:00 -> illegal right after P4

# Leave-usage order required when someone wants 3+ consecutive off days
LEAVE_PRIORITY = ["V", "HBD", "LK", "LP"]

DEFAULT_RULES = dict(
    max_consec_work=4, max_consec_plain_off=2,
    cap_C4=11, cap_CK=4, cap_P4=8,
    need_C4=3, need_P4=2, ck_counts_as_day=True,
)


def norm(v):
    """Normalise a cell value to a canonical code (or None if empty)."""
    if v is None:
        return None
    s = str(v).strip()
    if s == "":
        return None
    u = s.upper().replace(" ", "")
    if u in ("C4", "P4", "CK", "X", "V", "LP", "HBD", "LK", "H", "-"):
        return u
    if u == "C":                    # bare "C" = training/other day work (see comments e.g. อบรม)
        return "TRAIN"
    return "H?" + s                 # unknown code: keep as locked leave-like, warn later


def read_roster(path, sheet=None):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    days = []
    c = DAY_COL0
    while isinstance(ws.cell(HEADER_ROW, c).value, (int, float)):
        days.append(int(ws.cell(HEADER_ROW, c).value)); c += 1
    nurses = []
    r = FIRST_NURSE_ROW
    while isinstance(ws.cell(r, 1).value, (int, float)):
        grid = [norm(ws.cell(r, DAY_COL0 + d).value) for d in range(len(days))]
        nurses.append(dict(row=r, no=ws.cell(r, 1).value,
                           name=str(ws.cell(r, 3).value or "").strip(),
                           level=str(ws.cell(r, 5).value or "").strip(), grid=grid))
        r += 1
    return wb, ws, days, nurses


def is_off(code):
    return code is not None and (code in OFF or code.startswith("H?"))


def is_leave(code):
    return code is not None and (code in LEAVE or code.startswith("H?"))


def validate(nurses, ndays, rules=None, prev=None):
    """Return list of violation strings for a (possibly partial) roster.
    prev: {nurse_name: [codes of last days of previous month]} (optional)."""
    R = dict(DEFAULT_RULES, **(rules or {}))
    prev = prev or {}
    out = []
    for n in nurses:
        g = list(prev.get(n["name"], [])) + list(n["grid"])
        off0 = len(prev.get(n["name"], []))
        nm = n["name"] or f"row {n['row']}"
        label = lambda i: f"วันที่ {i - off0 + 1}"
        # consecutive work
        run = 0
        for i, c in enumerate(g):
            run = run + 1 if c in WORK else 0
            if run > R["max_consec_work"] and i >= off0:
                out.append(f"{nm}: ขึ้นเวรติดกันเกิน {R['max_consec_work']} วัน ({label(i)})")
        # reverse shift
        for i in range(len(g) - 1):
            if g[i] == "P4" and g[i + 1] in DAY_SHIFTS and i + 1 >= off0:
                out.append(f"{nm}: ขึ้นเวรย้อน P4→{g[i+1]} ({label(i)}→{label(i+1)})")
        # off run: day 3+ must be leave
        run = 0
        for i, c in enumerate(g):
            if c is None:
                run = 0; continue
            if is_off(c):
                run += 1
                if run >= 3 and i >= off0 and not is_leave(c):
                    out.append(f"{nm}: ออฟวันที่ 3+ ต่อเนื่องแต่ไม่ได้ระบุสิทธิ์การลา ({label(i)} = {c})")
            else:
                run = 0
        # monthly caps
        for code, cap in (("C4", R["cap_C4"]), ("CK", R["cap_CK"]), ("P4", R["cap_P4"])):
            k = sum(1 for c in n["grid"] if c == code)
            if k > cap:
                out.append(f"{nm}: {code} เกินโควตา {k}/{cap} วัน")
        for i, c in enumerate(n["grid"]):
            if c and c.startswith("H?"):
                out.append(f"{nm}: วันที่ {i+1} รหัส '{c[2:]}' ไม่รู้จัก (ถือเป็นวันลา)")
    # coverage
    for d in range(ndays):
        day = sum(1 for n in nurses if n["grid"][d] == "C4" or (R["ck_counts_as_day"] and n["grid"][d] == "CK"))
        ngt = sum(1 for n in nurses if n["grid"][d] == "P4")
        if day != R["need_C4"]:
            out.append(f"วันที่ {d+1}: เวรเช้า (C4{'+CK' if R['ck_counts_as_day'] else ''}) {day} คน ต้องการ {R['need_C4']}")
        if ngt != R["need_P4"]:
            out.append(f"วันที่ {d+1}: เวรดึก (P4) {ngt} คน ต้องการ {R['need_P4']}")
    return out
